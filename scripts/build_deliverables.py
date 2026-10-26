"""Genera borradores Word/PowerPoint a partir de fuentes revisables en Markdown."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DocInches
from docx.shared import Pt as DocPt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "deliverables"
IMAGE = re.compile(r"!\[(.*?)\]\((.*?)\)")


def architecture():
    """Diagrama determinista de dependencias; no es una captura ni evidencia de ejecución."""
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1280, 460), "#17232f")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("segoeui.ttf", 22)
    except OSError:
        font = ImageFont.load_default(size=22)

    def box(rect, text):
        draw.rounded_rectangle(rect, radius=12, fill="#244054", outline="#6bd7c5", width=2)
        draw.multiline_text((rect[0] + 16, rect[1] + 24), text, font=font, fill="white", spacing=10)

    def arrow(points):
        draw.line(points, fill="#6bd7c5", width=3)
        x, y = points[-1]
        draw.polygon([(x, y), (x - 10, y - 6), (x - 10, y + 6)], fill="#6bd7c5")

    box((20, 50, 235, 170), "GUI / Tk\nSolo presentación")
    box((285, 50, 530, 170), "Controlador\nWorker serial")
    box((580, 50, 815, 170), "VaultService\nSesión y casos")
    box((955, 10, 1260, 120), "Criptografía\nPBKDF2 + AES-GCM")
    box((955, 160, 1260, 270), "VaultRepository\nSolo sobres cifrados")
    box((955, 310, 1260, 420), "Dominio\nSnapshots inmutables")
    arrow([(235, 110), (285, 110)])
    arrow([(530, 110), (580, 110)])
    arrow([(815, 80), (890, 80), (955, 65)])
    arrow([(815, 110), (890, 110), (890, 215), (955, 215)])
    arrow([(815, 140), (860, 140), (860, 365), (955, 365)])
    box(
        (20, 280, 770, 420),
        "Laboratorio: worker independiente y datos ficticios\n"
        "Sin sesión, archivo ni credenciales de la bóveda",
    )
    image.save(DOCS / "evidence" / "architecture.png")


def paper():
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = DocPt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        document.styles[name].font.name = "Times New Roman"
    source = (DOCS / "PAPER_DRAFT.md").read_text(encoding="utf-8-sig")
    table_lines = []

    def flush_table():
        if not table_lines:
            return
        rows = [
            [cell.strip() for cell in row.strip("|").split("|")]
            for row in table_lines
            if not re.fullmatch(r"[| :\-]+", row)
        ]
        table = document.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        for row in rows:
            for cell, value in zip(table.add_row().cells, row, strict=True):
                cell.text = value.strip("`")
        table_lines.clear()

    for paragraph in source.split("\n\n"):
        text = " ".join(paragraph.splitlines()).strip()
        if not text:
            continue
        if paragraph.startswith("|"):
            table_lines.extend(paragraph.splitlines())
            flush_table()
        elif match := IMAGE.fullmatch(text):
            document.add_picture(str(DOCS / match[2]), width=DocInches(6))
            document.add_paragraph(match[1], style="Caption")
        elif text.startswith("#"):
            depth = len(text) - len(text.lstrip("#"))
            document.add_heading(text.lstrip("# "), level=0 if depth == 1 else depth - 1)
        else:
            document.add_paragraph(text.replace("**", "").replace("`", ""))
    document.save(OUTPUT / "VaultLab-paper-borrador.docx")


def slides():
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    for index, chunk in enumerate(
        (DOCS / "SLIDES_DRAFT.md").read_text(encoding="utf-8").split("\n---\n"), 1
    ):
        lines = [line.strip() for line in chunk.splitlines() if line.strip()]
        title = lines[0].lstrip("# ")
        body = [line for line in lines[1:] if not line.startswith("Notas:")]
        notes = [line.removeprefix("Notas:").strip() for line in lines if line.startswith("Notas:")]
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("17232F")

        def textbox(text, x, y, width, height, size, color="F0F4F8", target_slide=slide):
            shape = target_slide.shapes.add_textbox(
                Inches(x), Inches(y), Inches(width), Inches(height)
            )
            frame = shape.text_frame
            frame.word_wrap = True
            for number, line in enumerate(text.split("\n")):
                paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.size = Pt(size)
                paragraph.font.name = "Aptos"
                paragraph.font.color.rgb = RGBColor.from_string(color)
                paragraph.space_after = Pt(22)
            return shape

        textbox(title, 0.5, 0.25, 12.3, 0.8, 32, "6BD7C5")
        if body and (match := IMAGE.fullmatch(body[0])):
            from PIL import Image

            path = DOCS / match[2]
            with Image.open(path) as picture:
                ratio = picture.width / picture.height
            height = 5.5
            width = height * ratio
            if width > 11.8:
                width, height = 11.8, 11.8 / ratio
            slide.shapes.add_picture(
                str(path),
                Inches((13.333 - width) / 2),
                Inches(1.15),
                width=Inches(width),
                height=Inches(height),
            )
            textbox(match[1], 0.5, 6.75, 11.5, 0.4, 14)
        else:
            textbox("\n".join(body), 0.7, 1.6, 11.8, 5.0, 26)
        textbox(f"VaultLab · Académico · {index:02d}", 10.3, 7.0, 2.7, 0.3, 12)
        slide.notes_slide.notes_text_frame.text = "\n".join(notes)
    presentation.save(OUTPUT / "VaultLab-defensa-borrador.pptx")


def main():
    OUTPUT.mkdir(exist_ok=True)
    architecture()
    paper()
    slides()
    # Reabrir ambos archivos detecta estructura OOXML inválida, sin simular revisión Office.
    assert len(Document(OUTPUT / "VaultLab-paper-borrador.docx").paragraphs) > 20
    assert len(Presentation(OUTPUT / "VaultLab-defensa-borrador.pptx").slides) == 14


if __name__ == "__main__":
    main()
